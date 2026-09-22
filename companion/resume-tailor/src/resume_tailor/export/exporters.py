# Modified for the Career Agent public edition (2026-09-22). See NOTICE.
"""Export: one interface, several formats.

``Exporter.render(resume) -> bytes``. Markdown, HTML, DOCX and PDF (the DOCX
converted by a local Word or LibreOffice; a friendly error when neither is
installed). Every format renders the same ``GeneratedResume``, so the final
draft the user approved is what ships regardless of format. Adding a format
means adding one class and one registry entry.
"""

from __future__ import annotations

import html
import io
import re
from typing import Protocol

from resume_tailor.core.models import GeneratedResume
from resume_tailor.core.resume_generation.grouping import group_experience
from resume_tailor.core.text import format_ym


class Exporter(Protocol):
    content_type: str
    extension: str

    def render(self, resume: GeneratedResume) -> bytes: ...


def _dates(start: str, end: str | None) -> str:
    return f"{format_ym(start)} – {format_ym(end)}"


_INVALID_FS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def _sanitize(part: str) -> str:
    """A Windows-safe, readable filename fragment: invalid characters removed, whitespace
    collapsed, trailing periods/spaces trimmed, capitalization preserved."""
    part = _INVALID_FS.sub(" ", part)
    part = re.sub(r"\s+", " ", part).strip().rstrip(". ")
    return part


def export_filename(candidate_name: str, role_title: str, headline: str, extension: str) -> str:
    """User-facing artifact name: "<Applicant Name> - <Role Title>.<ext>".

    The role is the run's inferred JD title; an untitled or low-confidence role falls back to
    the recruiter-facing headline, and with neither the file is "<Applicant Name> - Resume".
    Run ids never appear in the visible name (internal run folders keep them for uniqueness)."""
    name = _sanitize(candidate_name) or "Resume"
    role = _sanitize(role_title) or _sanitize(headline) or "Resume"
    stem = f"{name} - {role}"[:120].rstrip(". ")
    return f"{stem}.{extension}"


def contact_parts(resume: GeneratedResume) -> list[tuple[str, str | None]]:
    """The header contact line as (visible text, link target or None). The visible
    text is exactly what the plain line showed, so linking changes no line width; the
    target adds ``mailto:`` / ``https://`` when the candidate wrote a bare address."""
    c = resume.candidate

    def web(v: str) -> str:
        return v if re.match(r"^[a-z][a-z0-9+.-]*://", v, re.IGNORECASE) else f"https://{v}"

    parts: list[tuple[str, str | None]] = []
    for value, kind in (
        (c.location, "text"),
        (c.phone, "text"),
        (c.email, "mail"),
        (c.linkedin, "web"),
        (c.portfolio, "web"),
    ):
        if not value:
            continue
        href = f"mailto:{value}" if kind == "mail" else web(value) if kind == "web" else None
        parts.append((value, href))
    return parts


def cert_lines(resume: GeneratedResume) -> list[str]:
    """One line per issuer: 'Workato: Automation Pro I, II and III; Advanced Data Transformations; ...'."""
    by_issuer: dict[str, list[str]] = {}
    for c in resume.certifications:
        by_issuer.setdefault(c.issuer, []).append(c.name)
    return [f"{issuer}: {'; '.join(names)}" for issuer, names in by_issuer.items()]


def _ordered_sections(resume: GeneratedResume) -> list[str]:
    default = ["summary", "experience", "skills", "education", "certifications"]
    order = [s for s in resume.section_order if s in default]
    return order + [s for s in default if s not in order]


class MarkdownExporter:
    content_type = "text/markdown; charset=utf-8"
    extension = "md"

    def render(self, r: GeneratedResume) -> bytes:
        c = r.candidate
        out = [
            f"# {c.name}",
            f"**{r.headline}**",
            "",
            " | ".join(f"[{text}]({href})" if href else text for text, href in contact_parts(r)),
            "",
        ]
        for s in _ordered_sections(r):
            if s == "summary" and r.summary:
                out += ["## Summary", " ".join(b.text for b in r.summary), ""]
            elif s == "experience":
                out.append("## Experience")
                for g in group_experience(r.experience):
                    if g.grouped:
                        # consecutive roles at one employer: company line once, every title and date kept
                        out.append(
                            f"**{g.company}**{', ' + g.location if g.location else ''} | {_dates(g.start, g.end)}  "
                        )
                        if g.company_blurb:
                            out.append(f"*{g.company_blurb}*  ")
                        for e in g.roles:
                            out.append(f"**{e.title}** | {_dates(e.start, e.end)}  ")
                            out += [f"- {b.text}" for b in e.bullets]
                        out.append("")
                        continue
                    e = g.roles[0]
                    out.append(f"**{e.title}**  ")
                    out.append(
                        f"{e.company}{', ' + e.location if e.location else ''} | {_dates(e.start, e.end)}  "
                    )
                    if e.company_blurb:
                        out.append(f"*{e.company_blurb}*  ")
                    out += [f"- {b.text}" for b in e.bullets]
                    out.append("")
            elif s == "skills" and r.skills:
                out.append("## Skills")
                out += [f"- **{g.name}:** {', '.join(g.items)}" for g in r.skills]
                out.append("")
            elif s == "education" and r.education:
                out.append("## Education")
                out += [
                    f"- {e.degree}, {e.institution}"
                    + (f" ({format_ym(e.start)} – {format_ym(e.end)})" if e.start else "")
                    for e in r.education
                ]
                out.append("")
            elif s == "certifications" and r.certifications:
                out.append("## Certifications")
                out += [f"- {line}" for line in cert_lines(r)]
                out.append("")
        if c.languages:
            out += ["## Languages", ", ".join(c.languages), ""]
        return "\n".join(out).encode("utf-8")


class HtmlExporter:
    content_type = "text/html; charset=utf-8"
    extension = "html"

    def render(self, r: GeneratedResume) -> bytes:
        c = r.candidate
        h = html.escape
        parts = [
            "<!doctype html><html><head><meta charset='utf-8'><title>",
            h(c.name),
            " - Resume</title>",
            "<style>body{font-family:Calibri,Arial,sans-serif;font-size:10.5pt;max-width:8.5in;margin:0.6in auto;color:#111;line-height:1.3}",
            "h1{font-size:18pt;margin:0}h2{font-size:11.5pt;text-transform:uppercase;border-bottom:1px solid #444;margin:14px 0 6px;letter-spacing:.5px}",
            ".hl{font-weight:bold;font-size:11.5pt}.meta{color:#333;margin:2px 0 8px}.role{font-weight:bold;margin-top:8px}.co{margin:0 0 3px}ul{margin:2px 0 6px 18px;padding:0}li{margin:2px 0}",
            ".blurb{font-style:italic;color:#444;margin:0 0 3px}.sub{font-weight:bold;margin-top:4px}",
            ".role,.sub,.co,.blurb{break-after:avoid;page-break-after:avoid}@media print{body{margin:0.4in}}</style></head><body>",
            f"<h1>{h(c.name)}</h1><div class='hl'>{h(r.headline)}</div>",
            "<div class='meta'>",
            " | ".join(
                f"<a href='{h(href)}'>{h(text)}</a>" if href else h(text)
                for text, href in contact_parts(r)
            ),
            "</div>",
        ]
        for s in _ordered_sections(r):
            if s == "summary" and r.summary:
                parts += ["<h2>Summary</h2><p>", h(" ".join(b.text for b in r.summary)), "</p>"]
            elif s == "experience":
                parts.append("<h2>Experience</h2>")
                for g in group_experience(r.experience):
                    if g.grouped:
                        parts.append(
                            f"<div class='role'>{h(g.company)}{(', ' + h(g.location)) if g.location else ''} &nbsp;|&nbsp; {h(_dates(g.start, g.end))}</div>"
                        )
                        if g.company_blurb:
                            parts.append(f"<div class='blurb'>{h(g.company_blurb)}</div>")
                        for e in g.roles:
                            parts.append(
                                f"<div class='sub'>{h(e.title)} &nbsp;|&nbsp; {h(_dates(e.start, e.end))}</div>"
                            )
                            parts.append(
                                "<ul>"
                                + "".join(f"<li>{h(b.text)}</li>" for b in e.bullets)
                                + "</ul>"
                            )
                        continue
                    e = g.roles[0]
                    parts.append(
                        f"<div class='role'>{h(e.title)}</div><div class='co'>{h(e.company)}{(', ' + h(e.location)) if e.location else ''} &nbsp;|&nbsp; {h(_dates(e.start, e.end))}</div>"
                    )
                    if e.company_blurb:
                        parts.append(f"<div class='blurb'>{h(e.company_blurb)}</div>")
                    parts.append(
                        "<ul>" + "".join(f"<li>{h(b.text)}</li>" for b in e.bullets) + "</ul>"
                    )
            elif s == "skills" and r.skills:
                parts.append(
                    "<h2>Skills</h2><ul>"
                    + "".join(
                        f"<li><b>{h(g.name)}:</b> {h(', '.join(g.items))}</li>" for g in r.skills
                    )
                    + "</ul>"
                )
            elif s == "education" and r.education:
                parts.append(
                    "<h2>Education</h2><ul>"
                    + "".join(
                        f"<li>{h(e.degree)}, {h(e.institution)}"
                        + (f" ({h(format_ym(e.start))} – {h(format_ym(e.end))})" if e.start else "")
                        + "</li>"
                        for e in r.education
                    )
                    + "</ul>"
                )
            elif s == "certifications" and r.certifications:
                parts.append(
                    "<h2>Certifications</h2><ul>"
                    + "".join(f"<li>{h(line)}</li>" for line in cert_lines(r))
                    + "</ul>"
                )
        if c.languages:
            parts.append("<h2>Languages</h2><p>" + h(", ".join(c.languages)) + "</p>")
        parts.append("</body></html>")
        return "".join(parts).encode("utf-8")


def _add_hyperlink(paragraph, text: str, href: str) -> None:
    """A clickable run: a proper external relationship plus a ``w:hyperlink`` wrapper,
    underlined in the paragraph's own colour. Word and LibreOffice keep it in PDF."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rid = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


class DocxExporter:
    content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extension = "docx"

    def render(self, r: GeneratedResume) -> bytes:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        doc = Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = Pt(40)
            section.left_margin = section.right_margin = Pt(48)
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        c = r.candidate
        p = doc.add_paragraph()
        run = p.add_run(c.name)
        run.bold = True
        run.font.size = Pt(17)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2 = doc.add_paragraph()
        p2.add_run(r.headline).bold = True
        contact = doc.add_paragraph()
        for i, (text, href) in enumerate(contact_parts(r)):
            if i:
                contact.add_run(" | ")
            if href:
                _add_hyperlink(contact, text, href)
            else:
                contact.add_run(text)

        def heading(text: str) -> None:
            hp = doc.add_paragraph()
            hr = hp.add_run(text.upper())
            hr.bold = True
            hr.font.size = Pt(11.5)
            hp.paragraph_format.space_before = Pt(10)
            hp.paragraph_format.space_after = Pt(2)
            hp.paragraph_format.keep_with_next = True

        for s in _ordered_sections(r):
            if s == "summary" and r.summary:
                heading("Summary")
                doc.add_paragraph(" ".join(b.text for b in r.summary))
            elif s == "experience":
                heading("Experience")

                def head_line(text: str, bold: bool = True, italic: bool = False, before: int = 0):
                    # heading-like lines stay with what follows them so a company or role
                    # heading is not stranded at the bottom of a page
                    hp = doc.add_paragraph()
                    run = hp.add_run(text)
                    run.bold = bold
                    run.italic = italic
                    hp.paragraph_format.space_after = Pt(0)
                    hp.paragraph_format.space_before = Pt(before)
                    hp.paragraph_format.keep_with_next = True
                    return hp

                for g in group_experience(r.experience):
                    if g.grouped:
                        head_line(
                            f"{g.company}{', ' + g.location if g.location else ''} | {_dates(g.start, g.end)}",
                            before=6,
                        )
                        if g.company_blurb:
                            head_line(g.company_blurb, bold=False, italic=True)
                        for e in g.roles:
                            head_line(f"{e.title} | {_dates(e.start, e.end)}", before=3)
                            for b in e.bullets:
                                bp2 = doc.add_paragraph(b.text, style="List Bullet")
                                bp2.paragraph_format.space_after = Pt(1)
                        continue
                    e = g.roles[0]
                    head_line(e.title, before=6)
                    head_line(
                        f"{e.company}{', ' + e.location if e.location else ''} | {_dates(e.start, e.end)}",
                        bold=False,
                    )
                    if e.company_blurb:
                        head_line(e.company_blurb, bold=False, italic=True)
                    for b in e.bullets:
                        bp2 = doc.add_paragraph(b.text, style="List Bullet")
                        bp2.paragraph_format.space_after = Pt(1)
            elif s == "skills" and r.skills:
                heading("Skills")
                for g in r.skills:
                    sp = doc.add_paragraph()
                    sp.add_run(f"{g.name}: ").bold = True
                    sp.add_run(", ".join(g.items))
                    sp.paragraph_format.space_after = Pt(1)
            elif s == "education" and r.education:
                heading("Education")
                for e in r.education:
                    doc.add_paragraph(
                        f"{e.degree}, {e.institution}"
                        + (f" ({format_ym(e.start)} – {format_ym(e.end)})" if e.start else "")
                    )
            elif s == "certifications" and r.certifications:
                heading("Certifications")
                for line in cert_lines(r):
                    cp2 = doc.add_paragraph(line)
                    cp2.paragraph_format.space_after = Pt(0)
        if c.languages:
            heading("Languages")
            doc.add_paragraph(", ".join(c.languages))
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()


class ExportUnavailableError(NotImplementedError):
    """This format cannot be produced on this computer right now; the message is
    written for the user and the other formats still work."""


PDF_UNAVAILABLE_MESSAGE = (
    "PDF export isn't available on this computer yet: it needs Microsoft Word "
    "or LibreOffice installed. Word and Markdown export still work."
)
PDF_FAILED_MESSAGE = (
    "PDF export didn't finish this time. Try again in a moment, or use Word or Markdown export."
)


class PdfExporter:
    """The PDF is the DOCX rendered by a local converter (Word, else LibreOffice), so
    all three final formats share one canonical draft and one visual structure, the
    text stays selectable/searchable and fonts are embedded by the renderer. Fully
    local: no cloud conversion, no uploads. Page counts can differ from Word's own
    measurement only by renderer metrics — the DOCX measurement stays canonical."""

    content_type = "application/pdf"
    extension = "pdf"

    def render(self, resume: GeneratedResume) -> bytes:
        from resume_tailor.export.pagination import PdfConversionError, docx_to_pdf

        try:
            out = docx_to_pdf(DocxExporter().render(resume))
        except PdfConversionError as e:
            raise ExportUnavailableError(PDF_FAILED_MESSAGE) from e
        if out is None:
            raise ExportUnavailableError(PDF_UNAVAILABLE_MESSAGE)
        return out[0]


EXPORTERS: dict[str, Exporter] = {
    "md": MarkdownExporter(),
    "html": HtmlExporter(),
    "docx": DocxExporter(),
    "pdf": PdfExporter(),
}
