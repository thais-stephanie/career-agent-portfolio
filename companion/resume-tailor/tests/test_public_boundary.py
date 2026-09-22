# Modified for the Career Agent public edition (2026-09-22). See NOTICE.
"""Public companion boundary: no cross-site access or automatic evidence handoff."""

import pytest
from fastapi.testclient import TestClient
from resume_tailor.api.app import create_app
from resume_tailor.providers.llm.factory import build_provider
from resume_tailor.workspace import WorkspaceStore
from resume_tailor.workspace.demo import create_demo_candidate
from resume_tailor.workspace.sources import add_source


def test_no_implicit_ai_provider():
    assert build_provider(env={}).info().provider == "none"
    assert build_provider(env={"OPENAI_API_KEY": "synthetic-not-a-key"}).info().provider == "none"
    with pytest.raises(ValueError, match="Unknown LLM_PROVIDER"):
        build_provider(env={"LLM_PROVIDER": "typo", "OPENAI_API_KEY": "synthetic-not-a-key"})


def test_http_boundary(tmp_path):
    client = TestClient(create_app(home=tmp_path), base_url="http://127.0.0.1:8766")
    assert client.get("/api/candidates").status_code == 200
    assert client.get("/api/candidates", headers={"Host": "evil.example"}).status_code == 403
    assert (
        client.get("/api/candidates", headers={"Origin": "https://evil.example"}).status_code == 403
    )
    assert (
        client.post(
            "/api/candidates",
            json={"name": "Synthetic"},
            headers={"Origin": "http://127.0.0.1:8765"},
        ).status_code
        == 403
    )
    assert client.post("/api/candidates", data={"name": "Synthetic"}).status_code == 415
    assert (
        client.post(
            "/api/candidates",
            json={"name": "Synthetic"},
            headers={"Origin": "http://127.0.0.1:8766"},
        ).status_code
        == 200
    )


def test_unreviewed_source_does_not_create_evidence(tmp_path):
    import io

    from docx import Document

    ws = create_demo_candidate(WorkspaceStore(tmp_path))
    before = ws.evidence_file.read_bytes()
    doc = Document()
    doc.add_heading("Experience", 1)
    doc.add_paragraph("Invented Person - Synthetic Employer")
    doc.add_paragraph("Administered Salesforce for ten years.")
    buf = io.BytesIO()
    doc.save(buf)
    add_source(ws, "synthetic.docx", buf.getvalue(), "resume")
    assert ws.evidence_file.read_bytes() == before
    assert "Administered Salesforce for ten years." not in str(ws.load_index().bank)
