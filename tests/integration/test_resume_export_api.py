"""Resume Workspace V2 PR 5: evidence ids at the boundary, and export.

Synthetic documents in temporary profile databases. The PDF tests print in
the installed Microsoft Edge (Chrome when Edge is absent) and SKIP, saying
so, on a machine with neither: a skipped PDF is never reported as verified.
"""

from __future__ import annotations

import io
import json
import tempfile
import threading
from collections.abc import Iterator
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest
from tests.integration.test_resume_editor_api import call, make_api
from tests.support_resume import HIDDEN_LINE, HIDDEN_PROJECT, confirm_cited, exportable

from career_agent.clock import new_id
from career_agent.domain.claims import VerifiedClaim
from career_agent.domain.enums import ClaimSource, ClaimType
from career_agent.resume_doc import export as exporter
from career_agent.resume_doc.ats import read_docx, read_pdf
from career_agent.resume_doc.models import ResumeDocument, upgrade_resume_document
from career_agent.resume_doc.store import ResumeStore, forget_resume_data
from career_agent.storage.db import connect, transaction
from career_agent.storage.repositories import ClaimRepo
from career_agent.storage.workspace_repo import ensure_candidate
from career_agent.web.api import JobsApi
from career_agent.web.server import ApiError

HAS_BROWSER = exporter.find_browser() is not None
needs_browser = pytest.mark.skipif(
    not HAS_BROWSER, reason="no Microsoft Edge or Chrome here: the PDF is NOT verified"
)


def _claim(conn: Any, key: str, *, verified: bool = True) -> None:
    ClaimRepo(conn).add(
        ensure_candidate(conn),
        VerifiedClaim(
            claim_key=key,
            claim_type=ClaimType.EMPLOYMENT,
            text=f"Synthetic {key}",
            source=ClaimSource.SELF_ATTESTED,
            verified=verified,
        ),
    )


@pytest.fixture
def api(tmp_path: Path) -> JobsApi:
    api = make_api(tmp_path, "a")
    with connect(api.config.db_path) as conn, transaction(conn):
        candidate = ensure_candidate(conn)
        _claim(conn, "k-ok")
        _claim(conn, "k-draft", verified=False)
        _claim(conn, "k-retired")
        repo = ClaimRepo(conn)
        current = repo.history(candidate, "k-retired")[-1]
        repo.supersede(candidate, current.next_revision(verified=False))
    return api


def _stored(api: JobsApi, doc: ResumeDocument) -> dict[str, Any]:
    with connect(api.config.db_path) as conn, transaction(conn):
        confirm_cited(conn, doc)
        ResumeStore(conn).create_document(doc)
    return call(api, "GET", f"/documents/{doc.id}")


def _summary(evidence: list[str], origin: str = "AI_REWRITE") -> dict[str, Any]:
    return {
        "id": new_id(),
        "text": "Connected billing to support.",
        "origin": origin,
        "evidence_ids": evidence,
    }


def _save(api: JobsApi, made: dict[str, Any], summary: dict[str, Any]) -> dict[str, Any]:
    doc = {**made["document"], "summary": summary}
    body = {"document": doc, "expected_sha256": made["sha256"]}
    return call(api, "PATCH", f"/documents/{made['id']}/working", body)


def _export(api: JobsApi, made: dict[str, Any], fmt: str, **extra: Any) -> dict[str, Any]:
    body = {"format": fmt, "expected_sha256": made["sha256"], **extra}
    return call(api, "POST", f"/documents/{made['id']}/exports", body)


# ------------------------------------------------------- evidence boundary


def _superseded_row_id(api: JobsApi) -> str:
    with connect(api.config.db_path) as conn:
        return str(
            conn.execute(
                "SELECT id FROM verified_claim WHERE claim_key = 'k-retired'"
                " AND superseded_by_id IS NOT NULL"
            ).fetchone()[0]
        )


@pytest.mark.parametrize(
    "evidence",
    ["k-missing", "k-draft", "k-retired", "k-other-profile", "superseded-row"],
)
def test_evidence_a_profile_has_not_confirmed_now_is_refused(
    api: JobsApi, tmp_path: Path, evidence: str
) -> None:
    if evidence == "k-other-profile":
        other = make_api(tmp_path, "b")
        with connect(other.config.db_path) as conn, transaction(conn):
            _claim(conn, "k-other-profile")
    if evidence == "superseded-row":
        evidence = _superseded_row_id(api)
    made = call(api, "POST", "/documents", {"title": "Mine"})
    summary = _summary([evidence])
    with pytest.raises(ApiError) as refused:
        _save(api, made, summary)
    assert refused.value.status == 400 and refused.value.code == "evidence_not_confirmed"
    assert refused.value.data == {"lines": [summary["id"]]}
    assert call(api, "GET", f"/documents/{made['id']}")["sha256"] == made["sha256"], "nothing saved"
    with pytest.raises(ApiError):
        call(api, "POST", "/documents", {"from": {**made["document"], "summary": summary}})


def test_a_blank_evidence_id_and_an_uncited_rewrite_are_invalid(api: JobsApi) -> None:
    made = call(api, "POST", "/documents", {"title": "Mine"})
    for summary in (_summary(["   "]), _summary([]), _summary([], "RULE_REWRITE")):
        with pytest.raises(ApiError) as refused:
            _save(api, made, summary)
        assert refused.value.status == 400


@pytest.mark.parametrize(
    "summary",
    [
        _summary(["k-ok"]),
        _summary(["k-ok"], "EVIDENCE_VERBATIM"),
        _summary([], "USER_AUTHORED"),
        _summary([], "IMPORTED"),
    ],
)
def test_confirmed_evidence_and_the_persons_own_lines_are_saved(
    api: JobsApi, summary: dict[str, Any]
) -> None:
    made = call(api, "POST", "/documents", {"title": "Mine"})
    _save(api, made, summary)
    saved = call(api, "GET", f"/documents/{made['id']}")["document"]["summary"]
    assert {k: saved[k] for k in summary} == summary


def test_an_edited_evidence_line_keeps_its_ids_only_while_they_hold(api: JobsApi) -> None:
    made = call(api, "POST", "/documents", {"title": "Mine"})
    line = {
        **_summary(["k-ok"], "EVIDENCE_VERBATIM"),
        "override": "EDITED",
        "original_text": "Connected billing.",
        "text": "Connected billing for 9 teams.",
    }
    _save(api, made, line)
    with pytest.raises(ApiError):
        _save(api, made, {**line, "evidence_ids": ["k-ok", "k-retired"]})


# ------------------------------------------------------------------ JSON


def test_json_is_the_revision_itself_hidden_content_included(api: JobsApi) -> None:
    made = _stored(api, exportable())
    out = _export(api, made, "JSON")
    assert out["verified"] and [c["check"] for c in out["checks"]] == ["ROUND_TRIP"]
    assert out["checks"][0]["status"] == "PASS"
    assert "file_path" not in out and str(api.config.db_path.parent) not in json.dumps(out)
    with connect(api.config.db_path) as conn:
        store = ResumeStore(conn)
        row = store.get_export(out["id"])
        revision = store.get_revision(row.revision_id)
        working_sha = store.get_document(made["id"]).working_sha256
        data = exporter.stored_file(conn, row).read_bytes()
    assert revision.content_sha256 == working_sha == made["sha256"]
    back = upgrade_resume_document(data.decode("utf-8"))
    assert back == revision.content
    payload = json.loads(data)
    assert set(payload) == set(ResumeDocument.model_fields) and payload["schema_version"] == "1.0"
    assert HIDDEN_LINE in data.decode("utf-8"), "JSON keeps hidden lines: it is the document"
    assert "Conceição".encode() in data, "accents as typed, in UTF-8"


def test_download_is_the_stored_file_under_a_professional_name(api: JobsApi) -> None:
    made = _stored(api, exportable())
    out = _export(api, made, "JSON")
    got = api.handle_api("GET", out["download"], {}, {})
    assert got.filename == "Morgan_Conceição_Exemplo_Resume.json"
    assert got.content_type == "application/json" and got.body.startswith(b"{")


# ------------------------------------------------------------------ DOCX


def test_docx_holds_the_visible_resume_in_order_in_real_styles(api: JobsApi) -> None:
    from docx import Document

    made = _stored(api, exportable())
    out = _export(api, made, "DOCX")
    assert out["verified"] and out["page_count"] is None
    status = {c["check"]: c["status"] for c in out["checks"]}
    assert status["PAGES"] == "NOT_MEASURED" and status["LAYOUT"] == "NOT_MEASURED"
    assert all(s in ("PASS", "NOT_MEASURED") for s in status.values()), status
    body = api.handle_api("GET", out["download"], {}, {}).body
    word = Document(io.BytesIO(body))
    styles = [(p.style.name, p.text) for p in word.paragraphs]
    assert styles[0] == ("Title", "Morgan Conceição Exemplo")
    assert ("Heading 1", "Experience") in styles
    assert (
        "List Bullet",
        "Mentorei a equipe de integração em São Paulo, reduzindo 35% do retrabalho.",
    ) in styles
    names = [s for s, _ in styles]
    assert {"Title", "Heading 1", "Heading 2", "List Bullet"} <= set(names)
    text = "\n".join(read_docx(body))
    assert HIDDEN_LINE not in text and HIDDEN_PROJECT not in text
    assert text.index("Experience") < text.index("Education") < text.index("Skills")
    assert not word.tables and not word.inline_shapes
    assert abs(word.sections[0].page_width.mm - 210) < 0.5


# ------------------------------------------------------------------- PDF


@needs_browser
def test_pdf_is_the_rendered_resume_with_a_clean_text_layer(api: JobsApi) -> None:
    from pypdf import PdfReader

    made = _stored(api, exportable())
    out = _export(api, made, "PDF", preview_pages=1)
    status = {c["check"]: c["status"] for c in out["checks"]}
    assert out["verified"] and status["PAGES"] == "PASS" and out["page_count"] == 1, out
    assert out["engine"] in ("msedge-headless", "chrome-headless")
    body = api.handle_api("GET", out["download"], {}, {}).body
    pages = read_pdf(body)
    text = "\n".join(pages)
    assert "Morgan Conceição Exemplo" in text and "São Paulo" in text
    assert HIDDEN_LINE not in text and HIDDEN_PROJECT not in text
    links = [
        a.get_object()["/A"]["/URI"]
        for page in PdfReader(io.BytesIO(body)).pages
        for a in page.get("/Annots") or []
    ]
    assert "https://portfolio.example/morgan" in links
    page = PdfReader(io.BytesIO(body)).pages[0]
    assert round(float(page.mediabox.width) / 72 * 25.4) == 210


@needs_browser
def test_a_pdf_that_disagrees_with_the_preview_is_kept_and_said_to_fail(api: JobsApi) -> None:
    made = _stored(api, exportable())
    out = _export(api, made, "PDF", preview_pages=3, preview_overflow=["experience/x"])
    pages = next(c for c in out["checks"] if c["check"] == "PAGES")
    layout = next(c for c in out["checks"] if c["check"] == "LAYOUT")
    assert pages["status"] == "FAIL" and pages["params"] == {"n": 1, "preview": 3}
    assert layout["status"] == "WARNING" and "TALLER_THAN_PAGE" in layout["issues"]
    assert out["verified"] is False, "a failed check is never a verified export"
    assert call(api, "GET", f"/documents/{made['id']}/exports")[0]["id"] == out["id"]


class _Hits(BaseHTTPRequestHandler):
    hits: list[str] = []

    def do_GET(self) -> None:  # noqa: N802
        _Hits.hits.append(self.path)
        self.send_response(200)
        self.end_headers()

    def log_message(self, *args: Any) -> None:
        pass


@pytest.fixture
def listener() -> Iterator[str]:
    server = ThreadingHTTPServer(("127.0.0.1", 0), _Hits)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    _Hits.hits = []
    try:
        yield f"http://127.0.0.1:{server.server_address[1]}"
    finally:
        server.shutdown()
        server.server_close()


@needs_browser
def test_printing_fetches_nothing_and_leaves_no_browser_profile(
    api: JobsApi, listener: str
) -> None:
    doc = exportable().model_dump(mode="json")
    doc["identity"]["links"].append({"id": new_id(), "kind": "OTHER", "url": f"{listener}/link"})
    doc["projects"][0]["url"] = f"{listener}/project"
    doc["layout"] = {}
    doc["summary"]["text"] = f'<img src="{listener}/img"> <link href="{listener}/css">'
    before = set(Path(tempfile.gettempdir()).glob("career-agent-pdf-*"))
    out = _export(api, _stored(api, upgrade_resume_document(doc)), "PDF")
    assert _Hits.hits == [], "rendering the resume visited an address"
    assert set(Path(tempfile.gettempdir()).glob("career-agent-pdf-*")) == before
    text = "\n".join(read_pdf(api.handle_api("GET", out["download"], {}, {}).body))
    assert "<img" in text, "markup typed into a field is text, never markup"


def test_no_browser_is_a_failed_file_never_a_checked_one(
    api: JobsApi, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(exporter, "find_browser", lambda: None)
    made = _stored(api, exportable())
    with pytest.raises(ApiError) as failed:
        _export(api, made, "PDF")
    assert failed.value.status == 503 and failed.value.code == "export_failed"
    assert call(api, "GET", f"/documents/{made['id']}/exports") == []


# --------------------------------------------------- refusals, revisions


def test_an_export_names_the_saved_copy_or_nothing_is_made(api: JobsApi) -> None:
    made = _stored(api, exportable())
    stale = {**made, "sha256": "0" * 64}
    with pytest.raises(ApiError) as refused:
        _export(api, stale, "JSON")
    assert refused.value.status == 409
    for body in (
        {"format": "PDF"},
        {"format": "TXT", "expected_sha256": made["sha256"]},
        {"format": "JSON", "expected_sha256": made["sha256"], "path": "C:/x.json"},
        {"format": "PDF", "expected_sha256": made["sha256"], "preview_pages": 0},
    ):
        with pytest.raises(ApiError) as bad:
            call(api, "POST", f"/documents/{made['id']}/exports", body)
        assert bad.value.status == 400
    assert call(api, "GET", f"/documents/{made['id']}/exports") == []


@pytest.mark.parametrize("name", ["", "   ", "You", "candidate"])
def test_no_real_name_no_file(api: JobsApi, name: str) -> None:
    doc = exportable().model_dump(mode="json")
    doc["identity"]["full_name"] = name
    made = _stored(api, upgrade_resume_document(doc))
    with pytest.raises(ApiError) as refused:
        _export(api, made, "DOCX")
    assert refused.value.code == "name_missing"


def test_evidence_retired_after_saving_stops_the_export(api: JobsApi) -> None:
    made = _stored(api, exportable())
    with connect(api.config.db_path) as conn, transaction(conn):
        candidate = ensure_candidate(conn)
        key = made["document"]["experience"][0]["bullets"][0]["evidence_ids"][0]
        repo = ClaimRepo(conn)
        repo.supersede(candidate, repo.history(candidate, key)[-1].next_revision(verified=False))
    with pytest.raises(ApiError) as refused:
        _export(api, made, "JSON")
    assert refused.value.code == "evidence_not_confirmed"
    assert refused.value.data == {"lines": [made["document"]["experience"][0]["bullets"][0]["id"]]}


def test_an_export_is_one_immutable_revision_and_changes_no_content(api: JobsApi) -> None:
    made = _stored(api, exportable())
    first = _export(api, made, "JSON")
    again = _export(api, made, "DOCX")
    with connect(api.config.db_path) as conn:
        store = ResumeStore(conn)
        a, b = store.get_export(first["id"]), store.get_export(again["id"])
        assert a.revision_id == b.revision_id, "the same saved copy is the same revision"
        assert len(store.list_revisions(made["id"])) == 1
    assert call(api, "GET", f"/documents/{made['id']}")["document"] == made["document"]
    edited = {**made["document"], "title": "Edited after the export"}
    call(
        api,
        "PATCH",
        f"/documents/{made['id']}/working",
        {"document": edited, "expected_sha256": made["sha256"]},
    )
    with connect(api.config.db_path) as conn:
        store = ResumeStore(conn)
        kept = store.get_revision(store.get_export(first["id"]).revision_id).content
    assert kept.title == made["document"]["title"]


def test_my_resumes_shows_the_last_export(api: JobsApi) -> None:
    made = _stored(api, exportable())
    _export(api, made, "JSON")
    out = _export(api, made, "DOCX")
    listed = next(d for d in call(api, "GET", "/documents") if d["id"] == made["id"])
    assert listed["last_export"]["id"] == out["id"] and listed["last_export"]["format"] == "DOCX"
    assert "score" not in json.dumps(listed).lower()


def test_a_name_full_of_path_characters_names_a_safe_file(api: JobsApi) -> None:
    doc = exportable().model_dump(mode="json")
    doc["identity"]["full_name"] = '..\\..//Ana:*?"<>| Lima\x07' + "x" * 200
    doc["title"] = "../../../outside"
    made = _stored(api, upgrade_resume_document(doc))
    out = _export(api, made, "JSON")
    name = api.handle_api("GET", out["download"], {}, {}).filename
    assert name.startswith("Ana_Lima") and name.endswith("_Resume.json") and len(name) < 90
    assert not set(name) & set('\\/:*?"<>|\x07')
    with connect(api.config.db_path) as conn:
        path = exporter.stored_file(conn, ResumeStore(conn).get_export(out["id"]))
    assert path.is_relative_to(api.config.db_path.parent / "resume_exports")


def test_forgetting_removes_the_exported_files(api: JobsApi) -> None:
    made = _stored(api, exportable())
    out = _export(api, made, "JSON")
    with connect(api.config.db_path) as conn:
        path = exporter.stored_file(conn, ResumeStore(conn).get_export(out["id"]))
        forget_resume_data(conn)
    assert not path.exists()


# ------------------------------------------------------- profile isolation


def test_one_profile_cannot_export_read_or_fetch_anothers(api: JobsApi, tmp_path: Path) -> None:
    other = make_api(tmp_path, "b")
    made = _stored(api, exportable())
    out = _export(api, made, "JSON")
    for method, path, body in (
        (
            "POST",
            f"/documents/{made['id']}/exports",
            {"format": "JSON", "expected_sha256": made["sha256"]},
        ),
        ("GET", f"/documents/{made['id']}/exports", {}),
    ):
        with pytest.raises(ApiError) as hidden:
            call(other, method, path, body)
        assert hidden.value.status == 404
    with pytest.raises(ApiError) as hidden:
        other.handle_api("GET", out["download"], {}, {})
    assert hidden.value.status == 404
    assert not (other.config.db_path.parent / "resume_exports").exists()
