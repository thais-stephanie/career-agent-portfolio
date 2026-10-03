"""Resume Workspace V2, PR 0: the current Resume helper stops lying.

Synthetic regression fixtures shaped like the failures met on a real install
(no owner data):

* a candidate still called by the placeholder "You" never reaches a file;
* a thin Career Profile (six roles, plain confirmed lines, no tenure record,
  no metrics or technologies) runs without inventing a richer model;
* a Portuguese ad flattened onto a few long lines yields requirements, each
  one the ad's own words, and an ad with no asks is said to be unreadable;
* a run keeps its target posting and the exact ad it read, and reopens so;
* a second run for the same job is a second version, the first untouched;
* the exported file is the draft on screen, or the export says why not.

No network, no model.
"""

from __future__ import annotations

import io
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from resume_tailor.core.job_analysis.analyzer import build_requirements, deterministic_analysis
from resume_tailor.core.models import TailorOptions, TailorRequest
from resume_tailor.core.pipeline import NoRequirementsFound, TailorService
from resume_tailor.export.exporters import EXPORTERS
from resume_tailor.integration.career import base_resume_from_profile, evidence_bank
from resume_tailor.providers.llm.vendors import NoneProvider
from resume_tailor.workspace import WorkspaceStore
from resume_tailor.workspace import drafts as dr
from tests.integration.test_tailor_bridge import TAILOR, _career_profile, _jobs, _profile

from career_agent.storage.db import connect

PT_TITLE = "Engenheiro(a) de Automação de IA - Sênior"
#: Shaped like the stored description that read as 0 requirements: a few very
#: long lines, Portuguese headings glued to the text, inline "•" items, one
#: semicolon list and narrative prose around them.
PT_JD = (
    "Você gosta de resolver problemas reais com automação? Na Exemplo Digital, construímos "
    "integrações para empresas de todos os tamanhos e acreditamos em times pequenos e "
    "autônomos que entregam com qualidade.Trabalhamos de forma remota, com ritmo sustentável e "
    "muita colaboração entre produto, dados e operações.Se você quer crescer com a gente, "
    "venha conversar.Responsabilidades e atribuições• Desenvolver integrações entre sistemas "
    "usando APIs REST e webhooks.• Automatizar processos internos com Python e n8n.• Garantir "
    "a observabilidade dos fluxos, com alertas e registros claros.• Documentar as soluções e "
    "apoiar o time de operações no dia a dia.\n"
    "Requisitos: Experiência com Python em produção; Conhecimento de APIs REST e webhooks; "
    "Vivência com ferramentas de automação como n8n ou Zapier; Inglês intermediário para "
    "leitura de documentação técnica; Capacidade de explicar soluções para pessoas não técnicas "
    "de forma clara e objetiva, sempre com cuidado com os detalhes.\n"
    "Diferenciais• Experiência com LangChain ou agentes de IA.• Conhecimento de SQL."
)

#: Six roles as Career Agent sends them: confirmed lines only, no metrics,
#: no technologies, no tenure record, and the placeholder name.
THIN_PROFILE = {
    "profile": {"id": "prof-01SYNTHETICTHINTHINTHINTH", "label": "My profile"},
    "name": "You",
    "experiences": [
        {
            "id": f"x{n}",
            "company": company,
            "title": title,
            "start": start,
            "end": end,
            "current": end is None,
            "kind": "EMPLOYMENT",
            "highlights": [
                {"key": f"k{n}a", "text": first, "origin": "document"},
                {"key": f"k{n}b", "text": second, "origin": "document"},
            ],
            "skills": [],
        }
        for n, (company, title, start, end, first, second) in enumerate(
            [
                (
                    "Northwind Ops",
                    "Automation Engineer",
                    "2023-01",
                    None,
                    "Built integrations between the billing system and the CRM.",
                    "Automated weekly reporting for the finance team.",
                ),
                (
                    "Contoso Health",
                    "Business Systems Analyst",
                    "2020-03",
                    "2022-12",
                    "Mapped recruiting workflows and documented the requirements.",
                    "Supported payroll operations during a system migration.",
                ),
                (
                    "Fabrikam Retail",
                    "Operations Analyst",
                    "2018-01",
                    "2020-02",
                    "Kept the inventory reports accurate for store managers.",
                    "Trained new analysts on the reporting process.",
                ),
                (
                    "Tailspin Logistics",
                    "Support Specialist",
                    "2016-05",
                    "2017-12",
                    "Answered customer questions about deliveries.",
                    "Wrote help articles for common delivery issues.",
                ),
                (
                    "Adventure Works",
                    "Office Assistant",
                    "2014-02",
                    "2016-04",
                    "Organised the office calendar and supplies.",
                    "Prepared monthly expense summaries.",
                ),
                (
                    "Wide World Importers",
                    "Intern",
                    "2013-01",
                    "2013-12",
                    "Entered shipping data into spreadsheets.",
                    "Helped the team check customs paperwork.",
                ),
            ]
        )
    ],
}


def _normal(text: str) -> str:
    return " ".join(text.split())


# =========================================================================
# the ad
# =========================================================================


def test_a_flattened_portuguese_ad_yields_requirements_in_its_own_words() -> None:
    job = build_requirements(deterministic_analysis(PT_JD))
    kinds = {r.category.value for r in job.requirements}
    assert len(job.requirements) >= 8, [r.text for r in job.requirements]
    assert {"must_have", "nice_to_have", "responsibility"} <= kinds
    ad = _normal(PT_JD)
    for requirement in job.requirements:
        if requirement.category.value != "technology":
            assert _normal(requirement.text) in ad, f"not the ad's words: {requirement.text!r}"
    # Prose about the company is not an ask.
    assert not any("Exemplo Digital" in r.text for r in job.requirements)


def test_an_ad_with_no_asks_is_unreadable_not_a_zero_score(tmp_path) -> None:
    ws, base = _thin_workspace(tmp_path)
    service = TailorService(ws.load_index(), ws.load_resumes(), ws.load_profiles(), NoneProvider())
    nothing = "Venha fazer parte de um time incrível! Somos uma empresa feliz. " * 4
    with pytest.raises(NoRequirementsFound):
        service.run(_request(nothing, base), "20260101T000000-aaaaaa")


# =========================================================================
# a thin profile
# =========================================================================


def _thin_workspace(tmp_path):  # noqa: ANN202
    store = WorkspaceStore(tmp_path / "home")
    ws = store.create("Synthetic Thin")
    base = base_resume_from_profile(ws, THIN_PROFILE)
    return ws, base


def _request(jd: str, base, **extra) -> TailorRequest:  # noqa: ANN001
    return TailorRequest(
        jd_text=jd, resume_id=base.id, options=TailorOptions(use_llm=False), **extra
    )


def test_the_placeholder_name_is_never_the_candidate() -> None:
    bank = evidence_bank(THIN_PROFILE, {"candidate": {"name": "You"}})
    assert bank["candidate"]["name"] == ""
    assert EXPORTERS["docx"].extension == "docx"
    from resume_tailor.export.exporters import export_filename

    assert not export_filename("You", "Engineer", "", "pdf").startswith("You")


def test_a_thin_profile_runs_for_the_target_without_rewriting_history(tmp_path) -> None:
    ws, base = _thin_workspace(tmp_path)
    service = TailorService(ws.load_index(), ws.load_resumes(), ws.load_profiles(), NoneProvider())
    run = service.run(
        _request(PT_JD, base, target_title=PT_TITLE, target_company="Exemplo Digital"),
        "20260101T000000-bbbbbb",
    )
    assert run.job_analysis.role_title == PT_TITLE
    assert run.job_analysis.company == "Exemplo Digital"
    assert run.job_analysis.requirements
    held = {e["title"] for e in THIN_PROFILE["experiences"]}
    # Historical titles are the person's own, never the target's.
    assert {e.title for e in run.generated_resume.experience} <= held
    # The headline is whole: the stable one when nothing safer can be said.
    assert run.generated_resume.headline in held | {PT_TITLE}
    assert run.generated_resume.candidate.name == ""
    # One chronology rule: the order the planner chose is the order the checks expect.
    assert not [i for i in run.validation_report.issues if i.code == "chronology"]
    assert not [f for f in run.lint_report.findings if f.code == "chronology"]


def test_the_exported_file_is_the_draft_on_screen(tmp_path) -> None:
    ws, base = _thin_workspace(tmp_path)
    service = TailorService(ws.load_index(), ws.load_resumes(), ws.load_profiles(), NoneProvider())
    run = service.run(_request(PT_JD, base, target_title=PT_TITLE), "20260101T000000-cccccc")
    index = ws.load_index()
    bullet = run.generated_resume.experience[0].bullets[0]
    shorter = bullet.text.rsplit(" ", 2)[0] + "."
    assert dr.validate_edit(run, index, bullet.id, shorter)["ok"], shorter
    state = {**dr.EMPTY_STATE, "bullets": {bullet.id: {"text": shorter}}}
    shown = dr.effective_resume(run, state, index)
    assert shown["experience"][0]["bullets"][0]["text"] == shorter
    exported = dr.export_resume(run, state, index, evidence_only=True)
    exported.candidate.name = "Riley Synthetic"
    text = "\n".join(
        p.text for p in Document(io.BytesIO(EXPORTERS["docx"].render(exported))).paragraphs
    )
    assert shorter in text
    assert bullet.text not in text, "the automatic wording shipped instead of the edit"
    # An edit the evidence does not support is named, never swapped.
    bold = {**dr.EMPTY_STATE, "bullets": {bullet.id: {"text": "Led a team of 40 engineers."}}}
    with pytest.raises(dr.ExportBlocked) as caught:
        dr.export_resume(run, bold, index, evidence_only=True)
    assert [line["text"] for line in caught.value.blocked] == ["Led a team of 40 engineers."]
    # With "only use things I have really done" off, the screen ships as shown.
    loose = dr.export_resume(run, bold, index, evidence_only=False)
    assert loose.experience[0].bullets[0].text == "Led a team of 40 engineers."


# =========================================================================
# through Career Agent: a run keeps its posting, versions, the name
# =========================================================================


@pytest.fixture
def helper(tmp_path):  # noqa: ANN201
    profile, api, client = _profile(tmp_path, "prof-01SYNTHETICPRZEROPRZEROPR", "Synthetic Zero")
    _career_profile(api)
    for path in ("/api/career/evidence/import", "/api/career/base-resume"):
        assert client.post(path, json={}, headers={"Origin": TAILOR}).status_code == 200
    cid = client.get("/api/workspace").json()["candidate_id"]
    return api, client, cid


def _run(client: TestClient, cid: str, job_id: str) -> str:
    started = client.post(
        f"/api/candidates/{cid}/tailor",
        json={"jd_text": "", "career_job_id": job_id, "options": {"use_llm": False}},
        headers={"Origin": TAILOR},
    )
    assert started.status_code == 200, started.text
    run_id = started.json()["application_id"]
    for _ in range(300):
        state = client.get(f"/api/candidates/{cid}/applications/{run_id}").json()["status"]
        if state["status"] in ("done", "error"):
            break
        time.sleep(0.2)
    assert state["status"] == "done", state
    return run_id


def _posting(api, job_id: str) -> tuple[str, str, str]:  # noqa: ANN001
    with connect(api.config.db_path) as conn:
        row = conn.execute(
            "SELECT j.title, c.name, r.description_text FROM job j JOIN company c"
            " ON c.id = j.company_id JOIN job_raw r ON r.content_hash = j.content_hash"
            " WHERE j.id = ?",
            (job_id,),
        ).fetchone()
    return str(row[0]), str(row[1]), str(row[2])


def test_a_run_reopens_with_its_posting_and_the_exact_ad(helper) -> None:
    api, client, cid = helper
    job_id = _jobs(api)[0]
    title, company, description = _posting(api, job_id)
    run_id = _run(client, cid, job_id)
    view = client.get(f"/api/candidates/{cid}/applications/{run_id}").json()["view"]
    assert view["job"]["role"] == title
    assert view["job"]["company"] == company
    assert view["job"]["ad"] == description
    assert view["job"]["source"]["career_job_id"] == job_id
    assert view["job"]["source"]["kind"] == "career_agent"
    assert view["assembled_without_ai"] is True


def test_a_second_run_is_a_second_version_and_the_first_is_untouched(helper) -> None:
    api, client, cid = helper
    job_id = _jobs(api)[0]
    first = _run(client, cid, job_id)
    before = client.get(f"/api/candidates/{cid}/applications/{first}").json()["view"]
    second = _run(client, cid, job_id)
    assert second != first
    runs = client.get(f"/api/candidates/{cid}/applications").json()
    assert sorted(r["id"] for r in runs) == sorted([first, second])
    assert {r["career_job_id"] for r in runs} == {job_id}
    assert client.get(f"/api/candidates/{cid}/applications/{first}").json()["view"] == before


def test_you_never_reaches_a_file_and_a_real_name_does(helper) -> None:
    api, client, cid = helper
    run_id = _run(client, cid, _jobs(api)[0])
    refused = client.get(f"/api/candidates/{cid}/applications/{run_id}/export/docx")
    assert refused.status_code == 400
    assert refused.json()["detail"]["code"] == "no_name"
    assert client.get("/api/workspace").json()["candidate_name"] == ""
    # The person gives a name in Career Agent; the next import carries it.
    api.handle_api("PATCH", "/api/candidate/name", {}, {"name": "Riley Synthetic"})
    assert client.post(
        "/api/career/evidence/import", json={}, headers={"Origin": TAILOR}
    ).is_success
    made = client.get(f"/api/candidates/{cid}/applications/{run_id}/export/docx")
    assert made.status_code == 200, made.text
    disposition = made.headers["content-disposition"]
    assert 'filename="Riley Synthetic - ' in disposition and "You" not in disposition
    text = "\n".join(p.text for p in Document(io.BytesIO(made.content)).paragraphs)
    assert "Riley Synthetic" in text
    assert client.get("/api/workspace").json()["candidate_name"] == "Riley Synthetic"
    # The placeholder itself cannot be given as a name.
    from career_agent.web.server import ApiError

    with pytest.raises(ApiError):
        api.handle_api("PATCH", "/api/candidate/name", {}, {"name": "You"})
