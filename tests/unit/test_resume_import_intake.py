"""Resume Workspace V2 PR 6: an uploaded resume is untrusted bytes.

Every refusal here is a code the page words; none quotes the file. Nothing
is executed (PDF JavaScript, DOCX macros), nothing is fetched (links,
external relationships), and a container cannot make the reader inflate
more than it declared. Synthetic files only.
"""

from __future__ import annotations

import io
import socket
import zipfile
from typing import Any

import pytest
from pypdf import PdfWriter
from tests.support_import import FOREIGN_HTML, IMAGE_ONLY_HTML, SIMPLE, docx, pdf_of

from career_agent.resume_doc import intake
from career_agent.resume_doc.export import find_browser
from career_agent.resume_doc.intake import ImportRefused, read_upload

needs_browser = pytest.mark.skipif(
    find_browser() is None, reason="no Edge or Chrome to print a PDF: NOT verified"
)


@pytest.fixture
def no_network(monkeypatch: pytest.MonkeyPatch) -> None:
    """Reading a resume must never open a connection."""

    def refuse(*_a: Any, **_k: Any) -> None:
        raise AssertionError("the reader tried to reach the network")

    monkeypatch.setattr(socket.socket, "connect", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)


def refused(data: bytes, name: str) -> str:
    with pytest.raises(ImportRefused) as caught:
        read_upload(data, name)
    return caught.value.code


def blank_pdf(pages: int = 1, **encrypt: str) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    if encrypt:
        writer.encrypt(**encrypt)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def rezip(data: bytes, extra: dict[str, bytes] | None = None, drop: tuple[str, ...] = ()) -> bytes:
    """A copy of a DOCX with members added, replaced or removed."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            if item.filename not in drop and item.filename not in (extra or {}):
                dst.writestr(item, src.read(item))
        for name, body in (extra or {}).items():
            dst.writestr(name, body, compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


# ------------------------------------------------------------------ the type


@pytest.mark.parametrize(
    "name", ["resume.doc", "resume.docm", "resume.rtf", "resume.png", "resume"]
)
def test_only_pdf_and_docx_are_accepted(name: str) -> None:
    assert refused(docx(SIMPLE), name) == "UNSUPPORTED_TYPE"


def test_the_bytes_must_be_what_the_name_says() -> None:
    assert refused(docx(SIMPLE), "resume.pdf") == "NOT_THE_TYPE"
    assert refused(blank_pdf(), "resume.docx") == "NOT_THE_TYPE"
    assert refused(b"plain text pretending", "resume.pdf") == "NOT_THE_TYPE"


def test_empty_and_oversized_files_are_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    assert refused(b"", "resume.pdf") == "EMPTY"
    monkeypatch.setattr(intake, "MAX_BYTES", 1000)
    assert refused(docx(SIMPLE), "resume.docx") == "TOO_LARGE"


# ------------------------------------------------------------------------ PDF


def test_a_password_protected_pdf_is_refused_by_name() -> None:
    assert (
        refused(blank_pdf(user_password="secret", owner_password="owner"), "r.pdf") == "ENCRYPTED"
    )


def test_too_many_pages_is_refused_before_reading_them() -> None:
    assert refused(blank_pdf(intake.MAX_PAGES + 1), "r.pdf") == "TOO_MANY_PAGES"


def test_a_pdf_without_text_is_said_so_never_an_empty_resume() -> None:
    assert refused(blank_pdf(), "r.pdf") == "NO_TEXT"


@pytest.mark.parametrize(
    "damage",
    [
        b"%PDF-1.7\n1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\ntrailer << /Root 1 0 R >>",
        b"%PDF-1.4\n" + b"\x00\xff" * 200,
    ],
)
def test_a_damaged_pdf_is_refused(damage: bytes) -> None:
    assert refused(damage, "r.pdf") in ("CORRUPT", "NO_TEXT")


@needs_browser
def test_an_image_only_pdf_is_no_text(no_network: None) -> None:
    assert refused(pdf_of(IMAGE_ONLY_HTML), "scan.pdf") == "NO_TEXT"


@needs_browser
def test_pdf_javascript_and_links_are_never_run_or_followed(no_network: None) -> None:
    writer = PdfWriter(clone_from=io.BytesIO(pdf_of(FOREIGN_HTML)))
    writer.add_js("app.alert('never');")
    writer.add_uri(0, "https://tracker.example.invalid/ping", [0, 0, 50, 50])
    out = io.BytesIO()
    writer.write(out)
    read = read_upload(out.getvalue(), "r.pdf")
    assert any(line.text.startswith("Avery Synthetic") for line in read.lines)


@needs_browser
def test_an_owner_password_only_pdf_opens_as_any_viewer_opens_it() -> None:
    writer = PdfWriter(clone_from=io.BytesIO(pdf_of(FOREIGN_HTML)))
    writer.encrypt(user_password="", owner_password="owner-only")
    out = io.BytesIO()
    writer.write(out)
    assert read_upload(out.getvalue(), "r.pdf").lines[0].text == "Avery Synthetic"


# ----------------------------------------------------------------------- DOCX


def test_a_malformed_zip_is_refused() -> None:
    assert refused(b"PK\x03\x04" + b"\x00" * 200, "r.docx") == "CORRUPT"


def test_a_zip_that_is_not_a_word_document_is_refused() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("readme.txt", "hello")
    assert refused(out.getvalue(), "r.docx") == "NOT_THE_TYPE"
    assert refused(rezip(docx(SIMPLE), drop=("word/document.xml",)), "r.docx") == "NOT_THE_TYPE"


def test_too_many_zip_entries_are_refused() -> None:
    extra = {f"word/media/pad{i}.txt": b"x" for i in range(intake.MAX_MEMBERS + 1)}
    assert refused(rezip(docx(SIMPLE), extra), "r.docx") == "UNSAFE_CONTAINER"


def test_a_decompression_bomb_is_refused_before_it_is_inflated() -> None:
    bomb = rezip(docx(SIMPLE), {"word/media/bomb.bin": b"\x00" * (25 * 1024 * 1024)})
    assert len(bomb) < 1_000_000, "small on disk"
    assert refused(bomb, "r.docx") == "UNSAFE_CONTAINER"


def test_macros_are_refused() -> None:
    data = docx(SIMPLE)
    assert refused(rezip(data, {"word/vbaProject.bin": b"macro"}), "r.docx") == "MACROS"
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        types = z.read("[Content_Types].xml").decode()
    docm = types.replace(
        "wordprocessingml.document.main+xml", "ms-word.document.macroEnabled.main+xml"
    )
    assert refused(rezip(data, {"[Content_Types].xml": docm.encode()}), "r.docx") == "MACROS"


def test_embedded_objects_and_external_relationships_are_never_opened(no_network: None) -> None:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    word = Document(io.BytesIO(docx(SIMPLE)))
    paragraph = word.paragraphs[0]
    rel = paragraph.part.relate_to(
        "https://tracker.example.invalid/x.png", RT.IMAGE, is_external=True
    )
    assert rel
    out = io.BytesIO()
    word.save(out)
    data = rezip(out.getvalue(), {"word/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0 object"})
    read = read_upload(data, "r.docx")
    assert "EMBEDDED_IGNORED" in read.warnings
    assert read.lines[0].text == "Morgan Example"


def test_a_docx_hyperlink_is_read_as_its_address() -> None:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from docx.oxml.parser import OxmlElement

    word = Document(io.BytesIO(docx(SIMPLE)))
    paragraph = word.add_paragraph()
    rid = paragraph.part.relate_to(
        "https://github.com/morgan-example", RT.HYPERLINK, is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "My code"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    out = io.BytesIO()
    word.save(out)
    read = read_upload(out.getvalue(), "r.docx")
    assert read.lines[-1].links == ("https://github.com/morgan-example",)


def test_text_and_the_document_body_are_bounded(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(intake, "MAX_TEXT", 200)
    assert "TEXT_CUT" in read_upload(docx(SIMPLE), "r.docx").warnings
    monkeypatch.setattr(intake, "MAX_BODY_BYTES", 100)
    assert refused(docx(SIMPLE), "r.docx") == "UNSAFE_CONTAINER"
