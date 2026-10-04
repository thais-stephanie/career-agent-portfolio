"""Synthetic resumes for the import tests (Resume Workspace V2 PR 6).

Every person, company and school here is invented. DOCX files are built with
python-docx in memory; foreign-layout PDFs are hand-written HTML (NOT the
Career Agent renderer) printed by the installed Edge or Chrome, with bullet
glyphs in the text and a company line above the role, the way word processors
lay resumes out.
"""

from __future__ import annotations

import io
from typing import Any

from career_agent.resume_doc.intake import SourceLine

EN, EM = chr(0x2013), chr(0x2014)  # dashes other people's documents contain

#: (kind, text) lines: "title", "h" (Heading 1), "h2" (Heading 2), "p", "b"
#: (bold paragraph), "li" (List Bullet), "table" (rows of cells: list[list[str]]).
Spec = list[tuple[str, Any]]


def docx(spec: Spec) -> bytes:
    from docx import Document
    from docx.shared import Pt

    word = Document()
    for kind, value in spec:
        if kind == "table":
            table = word.add_table(rows=len(value), cols=len(value[0]))
            for r, row in enumerate(value):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
            continue
        style = {
            "title": "Title",
            "h": "Heading 1",
            "h2": "Heading 2",
            "li": "List Bullet",
        }.get(kind, "Normal")
        paragraph = word.add_paragraph(style=style)
        run = paragraph.add_run(value)
        if kind == "b":
            run.bold = True
        if kind == "p":
            run.font.size = Pt(10.5)
    out = io.BytesIO()
    word.save(out)
    return out.getvalue()


SIMPLE: Spec = [
    ("title", "Morgan Example"),
    ("p", "morgan@example.invalid | +1 555 0100 | Springfield, US"),
    ("p", "https://www.linkedin.com/in/morgan-example-synthetic"),
    ("h", "Summary"),
    ("p", "Operations analyst who turns messy spreadsheets into reliable weekly reporting."),
    ("h", "Experience"),
    ("b", f"Operations Analyst | Northwind Traders | Jan 2021 {EN} Present"),
    ("li", "Rebuilt the weekly sales report so it runs in ten minutes instead of a day."),
    ("li", "Documented the returns process for the support team."),
    ("b", "Junior Analyst | Contoso Retail | 2018 - 2020"),
    ("li", "Reconciled store inventories every month."),
    ("h", "Education"),
    ("b", "BA, Economics | Example State University"),
    ("p", "2014 - 2018"),
    ("h", "Skills"),
    ("p", "Analysis: SQL, Excel, Revenue Operations"),
    ("p", "Languages: English, Spanish"),
]

#: The contact block is a two-column table, as some generators lay a header out.
TABLE_HEADER: Spec = [
    (
        "table",
        [["Riley Synthetic", "riley@example.invalid"], ["Data Engineer", "+44 20 7946 0000"]],
    ),
    ("h", "Experience"),
    ("b", "Data Engineer | Fabrikam Labs | 03/2019 - 08/2023"),
    ("li", "Moved nightly batch jobs to a streaming pipeline."),
    ("h", "Skills"),
    ("li", "Python"),
    ("li", "Event streams"),
]

UPPERCASE: Spec = [
    ("b", "Jordan Placeholder"),
    ("p", "jordan@example.invalid"),
    ("b", "PROFESSIONAL EXPERIENCE"),
    ("b", "Support Lead at Example Telecom"),
    ("p", "Jan 2020 to Dec 2022"),
    ("li", "Led a team of six support agents."),
    ("b", "EDUCATION"),
    ("b", "Example Technical College"),
    ("p", "2016"),
]

PORTUGUESE: Spec = [
    ("title", "Ana Exemplo Souza"),
    ("p", "ana.exemplo@example.invalid | (11) 99999-9999 | São Paulo, Brasil"),
    ("h", "Resumo Profissional"),
    ("p", "Analista de dados com foco em automação de relatórios."),
    ("h", "Experiência Profissional"),
    ("b", f"Analista de Dados | Exemplo Logística | jan 2021 {EN} atual"),
    ("li", "Automatizei o relatório semanal de entregas."),
    ("b", "Assistente Administrativa | Exemplo Comércio | 2018 - 2020"),
    ("li", "Organizei o controle de estoque da loja."),
    ("h", "Formação Acadêmica"),
    ("b", "Bacharelado em Administração | Universidade Exemplo"),
    ("p", "2014 - 2018"),
    ("h", "Competências"),
    ("p", "SQL, Power BI, Revenue Operations"),
    ("h", "Idiomas"),
    ("p", "Português (nativo), Inglês (avançado)"),
]

SPANISH: Spec = [
    ("title", "Lucía Ejemplo Ruiz"),
    ("p", "lucia@example.invalid | +34 600 000 000"),
    ("h", "Experiencia Profesional"),
    ("b", "Coordinadora de Proyectos | Ejemplo Consultores"),
    ("p", "ene 2020 - actualidad"),
    ("li", "Coordiné la migración de tres sistemas internos."),
    ("h", "Educación"),
    ("b", "Licenciatura en Economía | Universidad Ejemplo"),
    ("p", "2012 - 2016"),
    ("h", "Certificaciones"),
    ("p", "Gestión de Proyectos Ejemplo, Instituto Ejemplo, 2019"),
]

#: Early career: no summary, year-only dates, a custom section.
SPARSE: Spec = [
    ("title", "Sam Placeholder"),
    ("p", "sam@example.invalid"),
    ("h", "Education"),
    ("b", "BSc Computer Science | Example University"),
    ("p", "2021 - 2025"),
    ("h", "Projects"),
    ("b", "Campus bike map"),
    ("li", "Built a small map of bike racks for the student union."),
    ("h", "Volunteer Work"),
    ("b", "Volunteer | Example Food Bank | 2022"),
    ("li", "Sorted donations on weekends."),
    ("h", "Awards"),
    ("p", "Dean's list, 2023"),
]


def long_senior() -> Spec:
    spec: Spec = [
        ("title", "Casey Example"),
        ("p", "casey@example.invalid | +1 555 0199 | https://github.com/casey-example-synthetic"),
        ("h", "Summary"),
        ("p", "Engineering leader with fifteen years of building internal platforms."),
        ("h", "Experience"),
    ]
    for n in range(8):
        spec.append(("b", f"Engineering Manager | Example Company {n} | {2015 - n} - {2016 - n}"))
        spec += [("li", f"Delivered platform milestone {n}.{k} with the team.") for k in range(5)]
    spec += [("h", "Skills"), ("p", "Leadership: hiring, mentoring, planning")]
    return spec


#: Deliberately ambiguous: two prominent lines that could be the name, a
#: heading nobody uses, a header line whose second part is a place, and a
#: header with no role word on either side.
AMBIGUOUS: Spec = [
    ("b", "Taylor Sample"),
    ("b", "Jamie Sample"),
    ("p", "taylor@example.invalid"),
    ("h", "Experience"),
    ("b", "Data Analyst | Curitiba"),
    ("li", "Built dashboards."),
    ("b", "Example Bank | Example Insurance"),
    ("li", "Moved reports to the cloud."),
    ("h", "Things I Like"),
    ("p", "Chess and long walks."),
]

FOREIGN_HTML = f"""<!doctype html><meta charset="utf-8">
<style>
 body {{ font-family: Arial, sans-serif; font-size: 10.5pt; margin: 18mm; }}
 h1 {{ font-size: 22pt; margin: 0 0 4pt; }} h2 {{ font-size: 12pt; margin: 14pt 0 4pt; }}
 p {{ margin: 0 0 2pt; }} .b {{ font-weight: bold; }}
 .li {{ padding-left: 14pt; text-indent: -10pt; }}
</style>
<h1>Avery Synthetic</h1>
<p>avery@example.invalid &nbsp;|&nbsp; +55 21 98888-7777 &nbsp;|&nbsp; Rio de Janeiro, Brasil</p>
<h2>WORK EXPERIENCE</h2>
<p class="b">Example Airlines</p>
<p>Operations Planner {EM} Mar 2019 to Present</p>
<p class="li">• Planned weekly crew rosters for forty routes and kept them inside the rest rules
 that the regulator publishes every season, which used to take three people a full week.</p>
<p class="li">• Cut late changes by a third.</p>
<p class="b">Example Ground Services</p>
<p>Shift Coordinator {EM} 2016 to 2019</p>
<p class="li">• Coordinated ramp teams across two terminals.</p>
<h2>EDUCATION</h2>
<p class="b">Example Federal University</p>
<p>BSc in Logistics, 2012 {EN} 2016</p>
<h2>SKILLS</h2>
<p>Rostering, Excel, Revenue Operations, Stakeholder management</p>
"""

#: A page whose only content is a picture: no text layer at all.
IMAGE_ONLY_HTML = (
    '<!doctype html><img alt="" width="600" height="800" src="data:image/svg+xml;utf8,'
    "<svg xmlns='http://www.w3.org/2000/svg' width='600' height='800'>"
    "<rect width='600' height='800' fill='%23ddd'/><circle cx='300' cy='200' r='80'/></svg>\">"
)


def pdf_of(html: str) -> bytes:
    from career_agent.resume_doc.export import print_pdf

    return print_pdf(html)[0]


def lines(*texts: str, **kw: Any) -> list[SourceLine]:
    """SourceLines for parser unit tests: `!text` is bold, `#text` a heading
    style, `*text` a list item, `@text` a title style."""
    out = []
    for i, text in enumerate(texts, start=1):
        style, bold = "", False
        if text[:1] in "!#*@":
            style = {"#": "heading", "*": "list", "@": "title", "!": ""}[text[0]]
            bold = text[0] == "!"
            text = text[1:]
        out.append(SourceLine(text=text, where=f"paragraph {i}", bold=bold, style=style, **kw))
    return out
