"""Export one exact revision of a ResumeDocument as PDF, DOCX or JSON.

The file is always made from an immutable `resume_revision`, never from the
mutable working copy: `export_revision` checkpoints the working copy the
page has saved (reason EXPORTED; a copy equal to the latest revision IS that
revision, so nothing is duplicated), and everything after that reads only
that revision. The export row names it, and its `content_sha256` is the hash
of exactly what was exported.

* PDF prints the one renderer's print HTML in a headless Microsoft Edge
  (Chrome when Edge is absent): a throwaway profile in a temp directory, no
  window, no header or footer, every network address routed to a closed
  local port, and the document's own policy forbidding any load anyway.
* DOCX writes the SAME print HTML's blocks with real Word styles (Title,
  Heading 1, Heading 2, List Bullet, Normal). It never builds content of its
  own, so what is visible, in what order, with which dates, is the
  renderer's decision for both formats.
* JSON is the revision's ResumeDocument, schema version included.

Files go to the profile's private folder beside its database
(`resume_exports/<document>/<export>.<ext>`); the row stores the path
relative to that folder, and a download resolves it there or not at all.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import time
import unicodedata
from pathlib import Path
from typing import Any

from career_agent.clock import new_id
from career_agent.resume_doc.ats import blocks, check_export
from career_agent.resume_doc.evidence import unconfirmed_lines
from career_agent.resume_doc.models import ResumeDocument
from career_agent.resume_doc.render import ACCENTS, FONTS, PAGE_MM, render_html
from career_agent.resume_doc.store import (
    ExportFormat,
    ResumeExport,
    ResumeStore,
    StaleDocument,
)

EXTENSIONS: dict[str, str] = {"PDF": "pdf", "DOCX": "docx", "JSON": "json"}
CONTENT_TYPES = {
    "PDF": "application/pdf",
    "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "JSON": "application/json",
}
#: Seconds a PDF may take before the export is called failed.
PDF_TIMEOUT = 60


class ExportRefused(ValueError):
    """The export was not attempted: `code` says why, for the reader."""

    def __init__(self, code: str, lines: list[str] | None = None) -> None:
        super().__init__(code)
        self.code = code
        self.lines = lines or []


class ExportFailed(RuntimeError):
    """The file could not be made. Nothing was recorded."""


# ------------------------------------------------------------------ where


def export_dir(conn: sqlite3.Connection) -> Path:
    """This profile's private export folder: beside its database file."""
    for row in conn.execute("PRAGMA database_list"):
        if row[1] == "main" and row[2]:
            return Path(row[2]).parent / "resume_exports"
    raise ExportFailed("this database has no folder of its own")


def stored_file(conn: sqlite3.Connection, export: ResumeExport) -> Path:
    """The export's file, only if it lies inside this profile's export folder."""
    root = export_dir(conn).resolve()
    path = (root / export.file_path).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise FileNotFoundError(export.id)
    return path


_UNSAFE = re.compile(r'[<>:"/\\|?*\x00-\x1f\x7f]+')


def filename(doc: ResumeDocument, fmt: str) -> str:
    """`Full_Name_Resume.pdf`; a tailored version adds its job title. Letters
    stay as typed (NFC); characters no filesystem accepts become `_`."""

    def part(text: str) -> str:
        text = _UNSAFE.sub(" ", unicodedata.normalize("NFC", text))
        return "_".join(text.split()).strip("._")

    name = part(doc.identity.full_name)[:60].strip("._")
    title = part(doc.target.title)[:50].strip("._") if doc.target else ""
    return "_".join(p for p in (name, title, "Resume") if p) + "." + EXTENSIONS[fmt]


# --------------------------------------------------------------- writers


def find_browser() -> tuple[str, str] | None:
    """(path, engine) of Microsoft Edge, or Chrome when Edge is absent."""
    found: list[tuple[str, str]] = []
    if os.name == "nt":
        import winreg

        for exe, engine in (("msedge.exe", "msedge"), ("chrome.exe", "chrome")):
            for hive in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
                try:
                    key = rf"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\{exe}"
                    with winreg.OpenKey(hive, key) as handle:
                        found.append((winreg.QueryValue(handle, None), engine))
                except OSError:
                    pass
        for base in ("ProgramFiles(x86)", "ProgramFiles", "LOCALAPPDATA"):
            folder = os.environ.get(base) or ""
            found.append((str(Path(folder, "Microsoft/Edge/Application/msedge.exe")), "msedge"))
            found.append((str(Path(folder, "Google/Chrome/Application/chrome.exe")), "chrome"))
    for name, engine in (("msedge", "msedge"), ("google-chrome", "chrome"), ("chromium", "chrome")):
        which = shutil.which(name)
        if which:
            found.append((which, engine))
    found.sort(key=lambda f: f[1] != "msedge")  # Edge first, wherever it was found
    return next(((p, e) for p, e in found if p and Path(p).is_file()), None)


def print_pdf(html: str, *, timeout: float = PDF_TIMEOUT) -> tuple[bytes, str]:
    """Print the HTML to PDF in a headless browser. Returns (bytes, engine).

    Edge's launcher hands the work to a detached process and returns at
    once, so the PDF is awaited (written, of a stable size, and readable),
    then the temp profile is removed once the browser lets go of it."""
    from pypdf import PdfReader

    browser = find_browser()
    if browser is None:
        raise ExportFailed("NO_BROWSER")
    exe, engine = browser
    tmp = Path(tempfile.mkdtemp(prefix="career-agent-pdf-"))
    try:
        source, out = tmp / "resume.html", tmp / "resume.pdf"
        source.write_text(html, encoding="utf-8")
        subprocess.Popen(
            [
                exe,
                "--headless",
                "--disable-gpu",
                "--no-first-run",
                "--no-default-browser-check",
                "--disable-extensions",
                "--disable-sync",
                "--disable-background-networking",
                "--disable-component-update",
                f"--user-data-dir={tmp / 'profile'}",
                # Nothing may load from the network: every address goes to a
                # closed local port (the HTML's own policy forbids it as well).
                "--proxy-server=http://127.0.0.1:9",
                "--proxy-bypass-list=<-loopback>",
                "--no-pdf-header-footer",
                f"--print-to-pdf={out}",
                source.as_uri(),
            ],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        deadline, last = time.monotonic() + timeout, -1
        while time.monotonic() < deadline:
            size = out.stat().st_size if out.is_file() else 0
            if size and size == last:
                data = out.read_bytes()
                try:
                    PdfReader(io.BytesIO(data))
                    return data, f"{engine}-headless"
                except Exception:  # noqa: BLE001  -- still being written
                    pass
            last = size
            time.sleep(0.15)
        raise ExportFailed("PDF_TIMEOUT")
    finally:
        for _ in range(100):  # the browser may hold its profile a moment longer
            try:
                shutil.rmtree(tmp)
                break
            except FileNotFoundError:
                break
            except OSError:
                time.sleep(0.1)


def docx_bytes(doc: ResumeDocument, html: str) -> bytes:
    """A Word document of the rendered blocks, in real styles: no tables,
    text boxes, shapes or icons."""
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor

    design = doc.design
    width, height = PAGE_MM[design.page.size]
    font = FONTS[design.typography.font].split(",")[0].strip('"')
    accent = RGBColor.from_string(ACCENTS[design.accent].lstrip("#"))
    base = design.typography.base_pt
    word = Document()
    page = word.sections[0]
    page.page_width, page.page_height = Mm(width), Mm(height)
    margin = Mm(design.page.margins_mm)
    page.left_margin = page.right_margin = page.top_margin = page.bottom_margin = margin
    for name, size, color in (
        ("Normal", base, None),
        ("List Bullet", base, None),
        ("Title", base * 1.9, accent if design.template == "modern" else None),
        ("Heading 1", base * 1.05, accent),
        ("Heading 2", base, None),
    ):
        style = word.styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.color.rgb = color or RGBColor(0x1A, 0x1A, 0x1A)
        style.paragraph_format.line_spacing = design.typography.line_height
        style.paragraph_format.space_before = Pt(base * 0.6 if name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(2)
    styles = {
        "name": "Title",
        "heading": "Heading 1",
        "title": "Heading 2",
        "bullet": "List Bullet",
    }
    for block in blocks(html):
        paragraph = word.add_paragraph(style=styles.get(block.kind, "Normal"))
        for text, bold in block.runs:
            run = paragraph.add_run(text)
            run.bold = bold or block.kind == "headline" or None
            if block.kind == "meta":
                run.font.size = Pt(base * 0.92)
        # The print CSS's rules, said the way Word says them.
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = block.kind in ("heading", "title")
    props = word.core_properties
    props.title, props.author, props.comments = doc.title, doc.identity.full_name, ""
    props.last_modified_by = doc.identity.full_name
    out = io.BytesIO()
    word.save(out)
    return out.getvalue()


def json_bytes(doc: ResumeDocument) -> bytes:
    """The ResumeDocument itself: deterministic UTF-8, schema version included."""
    text = json.dumps(doc.model_dump(mode="json"), ensure_ascii=False, sort_keys=True, indent=2)
    return (text + "\n").encode("utf-8")


# ------------------------------------------------------------------ the act


def export_revision(
    conn: sqlite3.Connection,
    document_id: str,
    fmt: ExportFormat,
    *,
    expected_sha256: str,
    preview_pages: int | None = None,
    preview_overflow: list[str] | None = None,
    page_breaks: list[str] | None = None,
    pdf: Any = print_pdf,
) -> ResumeExport:
    """Export what the page has SAVED, as one revision, and check the file.

    Refused (nothing written) when the page holds a different copy, when a
    line cites evidence that is not confirmed now, or when there is no real
    name. `ExportFailed` when the file could not be made. A file that was
    made is kept and recorded even when a check finds a problem: the check
    result says exactly what. `page_breaks` are the blocks the preview
    started its pages with; the PDF starts its pages there too."""
    store = ResumeStore(conn)
    stored = store.get_document(document_id)
    if stored.working_sha256 != expected_sha256:
        raise StaleDocument(document_id, stored.working_sha256)
    lines = unconfirmed_lines(conn, stored.working)
    if lines:
        raise ExportRefused("EVIDENCE_NOT_CONFIRMED", lines)
    if stored.working.identity.name_finding():
        raise ExportRefused("NAME_MISSING")
    revision = store.checkpoint_revision(document_id, "EXPORTED")
    doc = revision.content
    # The PDF breaks pages where the preview did (`page_breaks`, its refs).
    html = render_html(doc, mode="print", breaks=frozenset(page_breaks or [])).html
    if fmt == "PDF":
        data, engine = pdf(html)
    elif fmt == "DOCX":
        data, engine = docx_bytes(doc, html), "python-docx"
    else:
        data, engine = json_bytes(doc), "json"
    report = check_export(
        fmt,
        doc,
        html,
        data,
        revision_sha256=revision.content_sha256,
        preview_pages=preview_pages,
        preview_overflow=preview_overflow or [],
    )
    export_id = new_id()
    relative = f"{document_id}/{export_id}.{EXTENSIONS[fmt]}"
    path = export_dir(conn) / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)

    return store.record_export(
        document_id,
        revision.id,
        format=fmt,
        template=doc.design.template,
        file_path=relative,
        file_sha256=hashlib.sha256(data).hexdigest(),
        engine=engine,
        page_count=report.get("pages"),
        ats_check=report,
        export_id=export_id,
    )
